import csv
import os
import re
import unicodedata

from openpyxl import Workbook, load_workbook

from src.database.eleve_dao import EleveDAO
from src.database.paiement_dao import PaiementDAO
from src.services.eleve_service import EleveService


class ImportExportService:
    COLONNES_ELEVES = [
        "Matricule",
        "Nom",
        "Prénom",
        "Classe",
        "Année scolaire",
        "Frais totaux",
    ]
    COLONNES_HISTORIQUE = [
        "Matricule",
        "Nom",
        "Prénom",
        "Classe",
        "N° reçu",
        "Date",
        "Montant",
        "Mode de règlement",
    ]

    @staticmethod
    def exporter_eleves(eleves, chemin):
        lignes = [
            [
                eleve.get("matricule", ""),
                eleve.get("nom", ""),
                eleve.get("prenom", ""),
                eleve.get("classe", ""),
                eleve.get("annee_scolaire", ""),
                eleve.get("montant_total_due", 0),
            ]
            for eleve in eleves
        ]
        ImportExportService._ecrire(chemin, ImportExportService.COLONNES_ELEVES, lignes)
        return chemin

    @staticmethod
    def exporter_historique(chemin):
        lignes = []
        for eleve in EleveDAO.obtenir_tous():
            for paiement in PaiementDAO.obtenir_par_eleve(eleve["id"]):
                lignes.append([
                    eleve.get("matricule", ""),
                    eleve["nom"],
                    eleve["prenom"],
                    eleve["classe"],
                    paiement.get("numero_recu", ""),
                    paiement.get("date_paiement", ""),
                    paiement.get("montant", 0),
                    paiement.get("mode_paiement", ""),
                ])
        ImportExportService._ecrire(chemin, ImportExportService.COLONNES_HISTORIQUE, lignes)
        return chemin

    @staticmethod
    def _ecrire(chemin, entetes, lignes):
        os.makedirs(os.path.dirname(os.path.abspath(chemin)), exist_ok=True)
        extension = os.path.splitext(chemin)[1].lower()
        if extension == ".csv":
            with open(chemin, "w", newline="", encoding="utf-8-sig") as fichier:
                writer = csv.writer(fichier, delimiter=";")
                writer.writerow(entetes)
                writer.writerows(lignes)
        elif extension == ".xlsx":
            classeur = Workbook()
            feuille = classeur.active
            feuille.title = "Élèves" if entetes == ImportExportService.COLONNES_ELEVES else "Historique"
            feuille.append(entetes)
            for ligne in lignes:
                feuille.append(ligne)
            feuille.freeze_panes = "A2"
            feuille.auto_filter.ref = feuille.dimensions
            for colonne in feuille.columns:
                longueur = max(len(str(cell.value or "")) for cell in colonne)
                feuille.column_dimensions[colonne[0].column_letter].width = min(longueur + 2, 32)
            classeur.save(chemin)
        else:
            raise ValueError("Le format de sortie doit être CSV ou XLSX.")

    @staticmethod
    def importer_eleves(chemin):
        lignes = ImportExportService._lire(chemin)
        if not lignes:
            return {"importes": 0, "erreurs": ["Le fichier est vide."]}

        index = {
            ImportExportService._normaliser(entete): position
            for position, entete in enumerate(lignes[0])
        }
        alias = {
            "matricule": ("matricule", "code eleve"),
            "nom": ("nom",),
            "prenom": ("prenom",),
            "classe": ("classe", "niveau"),
            "annee_scolaire": ("annee scolaire", "annee"),
            "montant_total_due": ("frais totaux", "montant total du", "scolarite"),
        }
        colonnes = {}
        for cle, choix in alias.items():
            colonnes[cle] = next((index[item] for item in choix if item in index), None)
        manquantes = [nom for nom in ("nom", "prenom", "classe", "annee_scolaire", "montant_total_due") if colonnes[nom] is None]
        if manquantes:
            return {"importes": 0, "erreurs": [f"Colonnes obligatoires absentes : {', '.join(manquantes)}."]}

        importes = 0
        erreurs = []
        for numero_ligne, ligne in enumerate(lignes[1:], start=2):
            try:
                valeurs = {
                    cle: (ligne[position] if position is not None and position < len(ligne) else "")
                    for cle, position in colonnes.items()
                }
                montant = float(str(valeurs["montant_total_due"]).replace(" ", "").replace(",", "."))
                EleveService.inscrire_eleve(
                    str(valeurs["nom"] or "").strip(),
                    str(valeurs["prenom"] or "").strip(),
                    str(valeurs["classe"] or "").strip(),
                    str(valeurs["annee_scolaire"] or "").strip(),
                    montant,
                    str(valeurs["matricule"] or "").strip() or None,
                )
                importes += 1
            except (ValueError, TypeError) as exc:
                erreurs.append(f"Ligne {numero_ligne} : {exc}")
        return {"importes": importes, "erreurs": erreurs}

    @staticmethod
    def _lire(chemin):
        extension = os.path.splitext(chemin)[1].lower()
        if extension == ".csv":
            with open(chemin, "r", newline="", encoding="utf-8-sig") as fichier:
                extrait = fichier.read(4096)
                fichier.seek(0)
                try:
                    dialecte = csv.Sniffer().sniff(extrait, delimiters=";,\t")
                    lecteur = csv.reader(fichier, dialecte)
                except csv.Error:
                    lecteur = csv.reader(fichier, delimiter=";")
                return list(lecteur)
        if extension == ".xlsx":
            classeur = load_workbook(chemin, read_only=True, data_only=True)
            try:
                return [list(ligne) for ligne in classeur.active.iter_rows(values_only=True)]
            finally:
                classeur.close()
        raise ValueError("Le fichier d'import doit être au format CSV ou XLSX.")

    @staticmethod
    def _normaliser(valeur):
        texte = unicodedata.normalize("NFKD", str(valeur or ""))
        texte = "".join(caractere for caractere in texte if not unicodedata.combining(caractere))
        return re.sub(r"[^a-z0-9]+", " ", texte.lower()).strip()
